"""
services/file_parser.py - Safe multi-format document parser for CareerMesh

Supports:
- Pasted text (with sanitization)
- .txt files
- .pdf files (via pypdf) with scanned-image detection
- .docx files (via python-docx)
"""

import os
from pathlib import Path
from typing import Tuple, Optional


MAX_FILE_SIZE_BYTES = 15 * 1024 * 1024  # 15 MB limit
SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}


class FileParsingError(Exception):
    """Custom exception raised when file cannot be parsed safely."""
    pass


def parse_text_input(raw_text: str) -> str:
    """Sanitizes raw text pasted by user."""
    if not raw_text or not raw_text.strip():
        raise FileParsingError("The provided text input is empty.")
    return raw_text.strip()


def parse_file(file_path: str | Path) -> Tuple[str, dict]:
    """
    Parses a file from disk (.txt, .pdf, .docx).

    Returns:
    (extracted_text, metadata_dict)

    Raises:
    FileParsingError if file is invalid, missing, oversized, or unreadable.
    """
    path = Path(file_path).resolve()

    if not path.exists():
        raise FileParsingError(f"File not found at path: {path}")

    if not path.is_file():
        raise FileParsingError(f"Provided path is not a file: {path}")

    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise FileParsingError(
            f"Unsupported file format '{ext}'. Supported formats: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    file_size = path.stat().st_size
    if file_size == 0:
        raise FileParsingError(f"File '{path.name}' is empty (0 bytes).")

    if file_size > MAX_FILE_SIZE_BYTES:
        raise FileParsingError(
            f"File size ({file_size / (1024*1024):.1f} MB) exceeds maximum limit of {MAX_FILE_SIZE_BYTES / (1024*1024):.0f} MB."
        )

    metadata = {
        "filename": path.name,
        "extension": ext,
        "size_bytes": file_size,
        "warning": None,
    }

    if ext == ".txt":
        return _parse_txt(path, metadata)
    elif ext == ".pdf":
        return _parse_pdf(path, metadata)
    elif ext == ".docx":
        return _parse_docx(path, metadata)
    else:
        raise FileParsingError(f"Handler for extension '{ext}' not found.")


def _parse_txt(path: Path, metadata: dict) -> Tuple[str, dict]:
    """Reads plain text file with encoding fallbacks."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            with open(path, "r", encoding=encoding) as f:
                content = f.read().strip()
                if not content:
                    raise FileParsingError(f"Text file '{path.name}' contains no readable text.")
                return content, metadata
        except UnicodeDecodeError:
            continue
        except Exception as e:
            raise FileParsingError(f"Error reading text file '{path.name}': {str(e)}")

    raise FileParsingError(f"Could not decode text file '{path.name}' using supported encodings.")


def _parse_pdf(path: Path, metadata: dict) -> Tuple[str, dict]:
    """Reads PDF pages using pypdf with scanned-page detection."""
    try:
        from pypdf import PdfReader
    except ImportError:
        raise FileParsingError(
            "pypdf is required to read PDF files. Please install it with 'pip install pypdf'."
        )

    try:
        reader = PdfReader(str(path))
        num_pages = len(reader.pages)
        metadata["num_pages"] = num_pages

        if num_pages == 0:
            raise FileParsingError(f"PDF '{path.name}' contains no pages.")

        extracted_chunks = []
        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            extracted_chunks.append(text)

        full_text = "\n".join(extracted_chunks).strip()

        # Check for scanned image PDF
        if len(full_text) < 30:
            warning_msg = (
                "Scanned or image-only PDF detected. The document contains very little or no selectable text. "
                "Optical Character Recognition (OCR) is required to parse images; no OCR was performed. "
                "Please copy and paste the text directly or provide a text-based PDF/DOCX."
            )
            metadata["warning"] = warning_msg
            if not full_text:
                raise FileParsingError(warning_msg)

        return full_text, metadata

    except FileParsingError:
        raise
    except Exception as e:
        raise FileParsingError(f"Failed to parse PDF '{path.name}': {str(e)}")


def _parse_docx(path: Path, metadata: dict) -> Tuple[str, dict]:
    """Reads DOCX paragraphs and tables using python-docx."""
    try:
        import docx
    except ImportError:
        raise FileParsingError(
            "python-docx is required to read DOCX files. Please install it with 'pip install python-docx'."
        )

    try:
        doc = docx.Document(str(path))
        text_lines = []

        # Extract text from paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                text_lines.append(para.text.strip())

        # Extract text from tables if any
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    text_lines.append(" | ".join(row_cells))

        full_text = "\n".join(text_lines).strip()
        if not full_text:
            raise FileParsingError(f"DOCX document '{path.name}' contains no readable text.")

        return full_text, metadata

    except FileParsingError:
        raise
    except Exception as e:
        raise FileParsingError(f"Failed to parse DOCX '{path.name}': {str(e)}")
